import os
import json
import tempfile
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware

from pypdf import PdfReader
from docx import Document

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client = Groq(api_key=my_api_key)
MODEL = "openai/gpt-oss-120b"

app = FastAPI(title="Resume Screener API")

# Allow the frontend (served from a different origin/port, or opened
# directly as a file) to call this API during local development.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------- Schemas ----------

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


class MatchResult(BaseModel):
    score: float
    details: dict


jobd_schema = JobD.model_json_schema()
resume_schema = Resume.model_json_schema()


# ---------- LLM calls ----------

def analyze_job_description(job_description: str) -> JobD:
    system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""
    user_prompt = f"""
Analyze the following job description:

{job_description}
"""
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return JobD(**data)


def parse_resume(resume_text: str) -> Resume:
    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning,
not only based on exact section headings.

Different resumes may use different headings.

For example:
- Experience
- Professional Experience
- Work History
- Employment
- Internships

These may all contain relevant experience.

Skills may also appear in the skills section, work experience,
internships or projects.

Return ONLY valid JSON matching this schema:

{resume_schema}

Important rules:

1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""
    user_prompt = f"""
Parse the following resume:

{resume_text}
"""
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return Resume(**data)


def final_score(job: JobD, resume: Resume) -> MatchResult:
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return JSON matching this schema:

{match_schema}

Give me:

1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict

Keep the response concise and easy to read.
"""
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)


# ---------- File reading ----------

def read_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_path: Path) -> str:
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path: Path) -> str | None:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix == ".docx":
        return read_docx(file_path)
    return None


# ---------- API endpoints ----------

@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.post("/api/screen")
async def screen_resumes(
    job_description: str = Form(...),
    resumes: list[UploadFile] = File(...),
):
    """
    Accepts a job description and one or more resume files.
    Returns every candidate ranked by match score, plus the
    parsed job requirements for reference.
    """
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description is empty.")
    if not resumes:
        raise HTTPException(status_code=400, detail="No resumes were uploaded.")

    job = analyze_job_description(job_description)

    results = []
    with tempfile.TemporaryDirectory() as tmp_dir:
        for upload in resumes:
            suffix = Path(upload.filename).suffix.lower()
            if suffix not in [".pdf", ".docx"]:
                results.append({
                    "filename": upload.filename,
                    "error": "Unsupported file type. Only .pdf and .docx are accepted.",
                })
                continue

            tmp_path = Path(tmp_dir) / upload.filename
            content = await upload.read()
            tmp_path.write_bytes(content)

            resume_text = read_resume(tmp_path)
            if not resume_text or not resume_text.strip():
                results.append({
                    "filename": upload.filename,
                    "error": "Could not extract any text from this file.",
                })
                continue

            try:
                parsed_resume = parse_resume(resume_text)
                match = final_score(job, parsed_resume)
                results.append({
                    "filename": upload.filename,
                    "name": parsed_resume.name,
                    "score": match.score,
                    "details": match.details,
                })
            except Exception as exc:
                results.append({
                    "filename": upload.filename,
                    "error": f"Failed to process this resume: {exc}",
                })

    # Rank the successfully-scored candidates; keep failures at the end.
    scored = [r for r in results if "score" in r]
    failed = [r for r in results if "score" not in r]
    scored.sort(key=lambda r: r["score"], reverse=True)

    return {
        "job": job.model_dump(),
        "candidates": scored + failed,
    }


if __name__ == "__main__":
    # Render (and most free hosts) inject the port to bind to via $PORT.
    # Running `python main.py` locally still works, defaulting to 8000.
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port)